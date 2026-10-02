"""In-place edits for Office documents: find and replace, and document properties.

Edits keep the document's formatting, styles, images, and layout. Modern
formats (DOCX, PPTX, XLSX) are edited directly; legacy and OpenDocument files
are first converted to the modern format of the same family. Full visual
editing is handed to the bundled Office engine's own editor.
"""
from __future__ import annotations

from pathlib import Path
import re
import subprocess
import sys
import tempfile
from typing import Callable

from office_engine import office_family, render_office
from studio_runtime import StagedOutput, check_cancel, find_libreoffice

_MODERN = {"writer": ".docx", "presentation": ".pptx", "spreadsheet": ".xlsx"}
PROPERTY_NAMES = ("title", "author", "subject", "keywords", "comments", "category")


def editable_extension(source: Path) -> str:
    """The format an edited copy of ``source`` is saved in."""
    family = office_family(source)
    if family not in _MODERN or source.suffix.lower() in {".csv", ".tsv"}:
        raise ValueError(f"{source.suffix or 'This file'} cannot be edited here. "
                         "Supported: Word, PowerPoint, and Excel family documents.")
    return _MODERN[family]


class _Editable:
    """Resolve ``source`` to a modern-format file, converting a copy when needed."""

    def __init__(self, source: Path, cancel_check):
        self.source, self.cancel_check = source, cancel_check
        self._directory = None

    def __enter__(self) -> Path:
        if not self.source.is_file():
            raise FileNotFoundError("The selected document no longer exists")
        extension = editable_extension(self.source)
        if self.source.suffix.lower() == extension:
            return self.source
        self._directory = tempfile.TemporaryDirectory(prefix="shadow-edit-")
        converted = Path(self._directory.name) / f"document{extension}"
        render_office(self.source, converted, extension.lstrip(".").upper(), self.cancel_check)
        return converted

    def __exit__(self, *_):
        if self._directory:
            self._directory.cleanup()


def _replace_in_runs(runs, pattern: re.Pattern, replacement: str) -> int:
    """Replace across a paragraph's runs, keeping each run's formatting.

    A match that spans several runs is written into the run where it starts.
    """
    texts = [run.text or "" for run in runs]
    joined = "".join(texts)
    matches = list(pattern.finditer(joined))
    if not matches:
        return 0
    bounds, position = [], 0
    for text in texts:
        bounds.append((position, position + len(text)))
        position += len(text)
    for match in reversed(matches):
        start, end = match.span()
        if start == end:
            continue
        first = next(i for i, (a, b) in enumerate(bounds) if a <= start < b)
        last = next(i for i, (a, b) in enumerate(bounds) if a < end <= b)
        if first == last:
            a = bounds[first][0]
            texts[first] = texts[first][:start - a] + match.expand(replacement) + texts[first][end - a:]
        else:
            texts[first] = texts[first][:start - bounds[first][0]] + match.expand(replacement)
            for middle in range(first + 1, last):
                texts[middle] = ""
            texts[last] = texts[last][end - bounds[last][0]:]
    for run, text in zip(runs, texts):
        if (run.text or "") != text:
            run.text = text
    return sum(1 for match in matches if match.start() != match.end())


def _docx_paragraphs(document):
    from docx.text.paragraph import Paragraph

    def walk(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from walk(part)
    # Text boxes, including the positioned frames of a PDF layout import.
    for box in document.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent"):
        for element in box.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
            yield Paragraph(element, document)


def _pptx_paragraphs(presentation):
    def shapes(collection):
        for shape in collection:
            if shape.shape_type == 6:  # group
                yield from shapes(shape.shapes)
            else:
                yield shape

    for slide in presentation.slides:
        containers = [slide.shapes]
        if slide.has_notes_slide:
            containers.append(slide.notes_slide.shapes)
        for container in containers:
            for shape in shapes(container):
                if shape.has_text_frame:
                    yield from shape.text_frame.paragraphs
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            yield from cell.text_frame.paragraphs


def replace_text(source: Path, destination: Path, replacements: dict[str, str], match_case: bool = False,
                 whole_word: bool = False, regex: bool = False, overwrite: bool = False,
                 cancel_check: Callable[[], bool] | None = None) -> tuple[Path, int]:
    """Find and replace in a document copy. Returns the output and the number of replacements."""
    if not replacements or not all(replacements):
        raise ValueError("Enter the text to find")
    extension = editable_extension(source)
    if destination.suffix.lower() != extension:
        raise ValueError(f"Save the edited copy as {extension}")
    if destination.resolve() == source.resolve():
        raise ValueError("Choose a new output path to keep your original document")
    compiled = []
    for find, replacement in replacements.items():
        body = find if regex else re.escape(find)
        if whole_word:
            body = rf"(?<!\w)(?:{body})(?!\w)"
        try:
            compiled.append((re.compile(body, 0 if match_case else re.I),
                             replacement if regex else replacement.replace("\\", "\\\\")))
        except re.error as exc:
            raise ValueError(f"Invalid pattern {find!r}: {exc}")
    count = 0
    with _Editable(source, cancel_check) as editable, StagedOutput(destination, overwrite) as stage:
        check_cancel(cancel_check)
        if extension == ".docx":
            from docx import Document
            document = Document(str(editable))
            for paragraph in _docx_paragraphs(document):
                for pattern, replacement in compiled:
                    count += _replace_in_runs(paragraph.runs, pattern, replacement)
            document.save(str(stage.path))
        elif extension == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(editable))
            for paragraph in _pptx_paragraphs(presentation):
                for pattern, replacement in compiled:
                    count += _replace_in_runs(paragraph.runs, pattern, replacement)
            presentation.save(str(stage.path))
        else:
            from openpyxl import load_workbook
            book = load_workbook(editable)
            try:
                for sheet in book.worksheets:
                    check_cancel(cancel_check)
                    for row in sheet.iter_rows():
                        for cell in row:
                            # Formulas are left alone; rewriting them could change results.
                            if isinstance(cell.value, str) and not cell.value.startswith("="):
                                value = cell.value
                                for pattern, replacement in compiled:
                                    value, made = pattern.subn(replacement, value)
                                    count += made
                                if value != cell.value:
                                    cell.value = value
                book.save(stage.path)
            finally:
                book.close()
        if count == 0:
            raise ValueError("The text was not found in this document")
    return stage.output, count


def read_properties(source: Path, cancel_check: Callable[[], bool] | None = None) -> dict[str, str]:
    """Title, author, subject, keywords, comments, and category."""
    extension = editable_extension(source)
    with _Editable(source, cancel_check) as editable:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(editable, read_only=True)
            try:
                core = book.properties
                return {"title": core.title or "", "author": core.creator or "", "subject": core.subject or "",
                        "keywords": core.keywords or "", "comments": core.description or "",
                        "category": core.category or ""}
            finally:
                book.close()
        if extension == ".docx":
            from docx import Document
            core = Document(str(editable)).core_properties
        else:
            from pptx import Presentation
            core = Presentation(str(editable)).core_properties
        return {name: getattr(core, name) or "" for name in PROPERTY_NAMES}


def write_properties(source: Path, destination: Path, properties: dict[str, str], overwrite: bool = False,
                     cancel_check: Callable[[], bool] | None = None) -> Path:
    """Save a copy with updated document properties. Unlisted properties are kept."""
    unknown = sorted(set(properties) - set(PROPERTY_NAMES))
    if unknown:
        raise ValueError(f"Unknown properties: {', '.join(unknown)}. Available: {', '.join(PROPERTY_NAMES)}")
    if not properties:
        raise ValueError("Enter at least one property to change")
    extension = editable_extension(source)
    if destination.suffix.lower() != extension:
        raise ValueError(f"Save the edited copy as {extension}")
    if destination.resolve() == source.resolve():
        raise ValueError("Choose a new output path to keep your original document")
    with _Editable(source, cancel_check) as editable, StagedOutput(destination, overwrite) as stage:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(editable)
            try:
                names = {"author": "creator", "comments": "description"}
                for name, value in properties.items():
                    setattr(book.properties, names.get(name, name), str(value))
                book.save(stage.path)
            finally:
                book.close()
        else:
            if extension == ".docx":
                from docx import Document
                document = Document(str(editable))
            else:
                from pptx import Presentation
                document = Presentation(str(editable))
            for name, value in properties.items():
                setattr(document.core_properties, name, str(value))
            document.save(str(stage.path))
    return stage.output


def open_in_editor(source: Path) -> None:
    """Open a document for full visual editing in the bundled Office engine."""
    if not source.is_file():
        raise FileNotFoundError("The selected document no longer exists")
    engine = find_libreoffice()
    if not engine:
        raise RuntimeError("Office engine is missing. Run setup_studio.ps1 or install LibreOffice.")
    executable = Path(engine)
    # soffice.com is the console launcher; the editor window belongs to soffice.exe.
    if executable.name.lower() == "soffice.com" and executable.with_suffix(".exe").is_file():
        executable = executable.with_suffix(".exe")
    flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    subprocess.Popen([str(executable), "--nologo", "--norestore", str(source.resolve())],
                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, creationflags=flags)
