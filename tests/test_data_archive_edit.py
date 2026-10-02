from __future__ import annotations

import io
import json
import os
from pathlib import Path
import sys
import tarfile
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import _headless  # noqa: F401

from studio_runtime import StudioCancelled


class DataConversionTests(unittest.TestCase):
    def test_csv_json_yaml_xml_xlsx_roundtrip(self):
        from data_tools import convert_data, load_records
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "leads.csv"
            source.write_text("Name,Budget,Active,Postcode\nAcme,1200,true,01234\nZed Ltd,80.5,false,\n", encoding="utf-8")
            records = load_records(source)
            self.assertEqual(records[0], {"Name": "Acme", "Budget": 1200, "Active": True, "Postcode": "01234"})
            self.assertIsNone(records[1]["Postcode"])
            for extension in (".json", ".jsonl", ".yaml", ".xml", ".xlsx", ".tsv"):
                output, count = convert_data(source, root / f"leads{extension}")
                self.assertEqual(count, 2)
                back = load_records(output)
                self.assertEqual(len(back), 2, extension)
                self.assertEqual(str(back[0]["Name"]), "Acme", extension)
                self.assertEqual(float(back[1]["Budget"]), 80.5, extension)
            with self.assertRaises(ValueError):
                convert_data(source, source)
            with self.assertRaises(ValueError):
                convert_data(source, root / "leads.docx")

    def test_nested_json_flattens_to_columns_and_semicolon_csv_is_detected(self):
        from data_tools import convert_data, load_records
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "campaigns.json"
            source.write_text(json.dumps({"campaigns": [
                {"name": "Spring", "owner": {"team": "Paid"}, "tags": ["a", "b"]},
                {"name": "Autumn", "owner": {"team": "Brand"}, "tags": []}]}), encoding="utf-8")
            output, _ = convert_data(source, root / "campaigns.csv")
            rows = load_records(output)
            self.assertEqual(rows[0]["owner.team"], "Paid")
            self.assertEqual(rows[0]["tags"], "a; b")
            semicolon = root / "eu.csv"
            semicolon.write_text("a;b\n1;2\n", encoding="utf-8")
            self.assertEqual(load_records(semicolon), [{"a": 1, "b": 2}])

    def test_unsafe_yaml_and_xml_are_rejected(self):
        from data_tools import load_records
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "evil.yaml").write_text("!!python/object/apply:os.system ['echo hi']", encoding="utf-8")
            with self.assertRaises(ValueError):
                load_records(root / "evil.yaml")
            (root / "bomb.xml").write_text(
                '<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]><d><r>&b;</r></d>',
                encoding="utf-8")
            with self.assertRaises(ValueError):
                load_records(root / "bomb.xml")

    def test_cleanup_reports_each_change(self):
        from data_tools import clean_data, load_records
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "messy.csv"
            source.write_text("First Name , Notes,Empty\n  Ann  ,ok,\n,,\nAnn,ok,\nBob,  two   spaces ,\n", encoding="utf-8")
            output, report = clean_data(source, root / "clean.csv", drop_duplicates=True, normalize_headers=True)
            self.assertEqual(report["empty_rows_removed"], 1)
            self.assertEqual(report["empty_columns_removed"], 1)
            self.assertEqual(report["duplicate_rows_removed"], 1)
            self.assertEqual(report["rows_written"], 2)
            rows = load_records(output, infer_types=False)
            self.assertEqual(list(rows[0]), ["first_name", "notes"])
            self.assertEqual(rows[1]["notes"], "two spaces")
            self.assertTrue(source.exists())

    def test_schema_validation_and_inference(self):
        from data_tools import infer_schema, load_records, validate_data
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "budget.csv"
            source.write_text("channel,budget\nSearch,1200\nSocial,lots\n,5\n", encoding="utf-8")
            schema = {"type": "object", "required": ["channel", "budget"],
                      "properties": {"channel": {"type": "string"}, "budget": {"type": "number", "minimum": 0}}}
            problems = validate_data(source, schema)
            self.assertEqual({(p["record"], p["field"]) for p in problems}, {(2, "budget"), (3, "channel")})
            good = root / "good.json"
            good.write_text(json.dumps([{"channel": "Search", "budget": 5}]), encoding="utf-8")
            self.assertEqual(validate_data(good, schema), [])
            inferred = infer_schema(load_records(good))
            self.assertEqual(inferred["items"]["properties"]["budget"], {"type": "integer"})
            self.assertEqual(validate_data(good, inferred), [])
            with self.assertRaises(ValueError):
                validate_data(good, {"type": "nonsense"})

    def test_every_sheet_is_exported(self):
        from openpyxl import Workbook
        from data_tools import export_sheets, sheet_names
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            book = Workbook()
            book.active.title = "Q1"
            book.active.append(["Channel", "Spend"])
            book.active.append(["Search", 10])
            second = book.create_sheet("Q2 Plan")
            second.append(["Channel", "Spend"])
            second.append(["Social", 20])
            book.create_sheet("Empty")
            book.save(root / "plan.xlsx")
            book.close()
            self.assertEqual(sheet_names(root / "plan.xlsx"), ["Q1", "Q2 Plan", "Empty"])
            outputs = export_sheets(root / "plan.xlsx", root / "sheets")
            self.assertEqual([p.name for p in outputs], ["plan-Q1.csv", "plan-Q2 Plan.csv"])
            self.assertIn("Social,20", outputs[1].read_text(encoding="utf-8-sig"))
            with self.assertRaises(ValueError):
                export_sheets(root / "plan.xlsx", root / "sheets", "XLSX")


class ArchiveTests(unittest.TestCase):
    def _tree(self, root: Path) -> Path:
        folder = root / "deliverables"
        (folder / "images").mkdir(parents=True)
        (folder / "report.txt").write_text("final report", encoding="utf-8")
        (folder / "images" / "hero.bin").write_bytes(b"\x00\x01" * 500)
        return folder

    def test_create_inspect_extract_every_format(self):
        from archive_tools import create_archive, extract_archive, inspect_archive
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            folder = self._tree(root)
            for name in ("pack.zip", "pack.7z", "pack.tar.gz", "pack.tar.xz", "pack.tar"):
                archive = create_archive([folder], root / name)
                report = inspect_archive(archive)
                self.assertEqual(report["files"], 2, name)
                self.assertEqual(report["problems"], [], name)
                out = extract_archive(archive, root / "out")
                self.assertEqual(out.name, "pack" if name == "pack.zip" else out.name)
                self.assertEqual((out / "deliverables" / "report.txt").read_text(encoding="utf-8"), "final report")
                self.assertEqual((out / "deliverables" / "images" / "hero.bin").stat().st_size, 1000)
            self.assertEqual(sorted(p.name for p in (root / "out").iterdir()),
                             ["pack", "pack_1", "pack_2", "pack_3", "pack_4"])

    def test_encrypted_7z_needs_its_password(self):
        from archive_tools import create_archive, extract_archive
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            archive = create_archive([self._tree(root)], root / "secret.7z", password="hunter2")
            with self.assertRaises(ValueError):
                extract_archive(archive, root / "out")
            with self.assertRaises(ValueError):
                extract_archive(archive, root / "out", password="wrong")
            out = extract_archive(archive, root / "out", password="hunter2")
            self.assertTrue((out / "deliverables" / "report.txt").is_file())
            with self.assertRaises(ValueError):
                create_archive([root / "deliverables"], root / "secret.zip", password="x")
            self.assertEqual([p.name for p in (root / "out").iterdir()], [out.name])

    def test_traversal_links_and_bombs_are_refused(self):
        from archive_tools import UnsafeArchive, extract_archive, inspect_archive
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            evil = root / "evil.zip"
            with zipfile.ZipFile(evil, "w") as archive:
                archive.writestr("../escape.txt", "x")
            self.assertTrue(inspect_archive(evil)["problems"])
            with self.assertRaises(UnsafeArchive):
                extract_archive(evil, root / "out")
            absolute = root / "absolute.zip"
            with zipfile.ZipFile(absolute, "w") as archive:
                archive.writestr("C:/Windows/evil.txt", "x")
            with self.assertRaises(UnsafeArchive):
                extract_archive(absolute, root / "out")
            linked = root / "link.tar"
            with tarfile.open(linked, "w") as archive:
                info = tarfile.TarInfo("link")
                info.type = tarfile.SYMTYPE
                info.linkname = "/etc/passwd"
                archive.addfile(info)
            with self.assertRaises(UnsafeArchive):
                extract_archive(linked, root / "out")
            big = root / "big.zip"
            with zipfile.ZipFile(big, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("zeros.bin", b"\x00" * 2_000_000)
            with self.assertRaises(UnsafeArchive):
                extract_archive(big, root / "out", max_bytes=1_000_000)
            with self.assertRaises(UnsafeArchive):
                extract_archive(big, root / "out", max_entries=0)
            self.assertFalse((root / "escape.txt").exists())
            self.assertFalse(any((root / "out").iterdir()) if (root / "out").exists() else False)
            junk = root / "junk.zip"
            junk.write_bytes(b"not an archive")
            with self.assertRaises(ValueError):
                inspect_archive(junk)

    def test_cancel_leaves_no_partial_folder(self):
        from archive_tools import create_archive, extract_archive
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            archive = create_archive([self._tree(root)], root / "pack.zip")
            with self.assertRaises(StudioCancelled):
                extract_archive(archive, root / "out", cancel_check=lambda: True)
            self.assertEqual(list((root / "out").iterdir()), [])

    def test_manifest_detects_changed_missing_and_unlisted_files(self):
        from archive_tools import CHECKSUM_NAME, verify_folder, write_folder_manifest
        with tempfile.TemporaryDirectory() as td:
            folder = self._tree(Path(td))
            manifest = write_folder_manifest(folder, {"client": "Acme"})
            data = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(data["file_count"], 2)
            self.assertEqual(data["client"], "Acme")
            self.assertIn("*report.txt", (folder / CHECKSUM_NAME).read_text(encoding="utf-8"))
            self.assertTrue(verify_folder(folder)["ok"])
            (folder / "report.txt").write_text("edited report", encoding="utf-8")
            (folder / "images" / "hero.bin").unlink()
            (folder / "new.txt").write_text("surprise", encoding="utf-8")
            report = verify_folder(folder)
            self.assertFalse(report["ok"])
            self.assertEqual((report["changed"], report["missing"], report["unlisted"]),
                             (["report.txt"], ["images/hero.bin"], ["new.txt"]))

    def test_delivery_package_verifies_and_detects_tampering(self):
        from archive_tools import MANIFEST_NAME, build_delivery_package, verify_package
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            folder = self._tree(root)
            package = build_delivery_package([folder], root / "acme-delivery.zip", client="Acme", project="Launch",
                                             notes="Final assets")
            with zipfile.ZipFile(package) as archive:
                names = set(archive.namelist())
                self.assertTrue({MANIFEST_NAME, "SHA256SUMS.txt", "README.txt"} <= names)
                self.assertIn("Acme", archive.read("README.txt").decode("utf-8"))
            report = verify_package(package)
            self.assertTrue(report["ok"], report)
            self.assertEqual(report["checked"], 2)
            tampered = root / "tampered.zip"
            with zipfile.ZipFile(package) as original, zipfile.ZipFile(tampered, "w") as copy:
                for info in original.infolist():
                    data = original.read(info)
                    copy.writestr(info, b"changed report" if info.filename.endswith("report.txt") else data)
            self.assertEqual(verify_package(tampered)["changed"], ["deliverables/report.txt"])


class OfficeEditTests(unittest.TestCase):
    def test_docx_replace_keeps_formatting_across_runs_tables_and_headers(self):
        from docx import Document
        from office_edit import replace_text
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Prepared for Old").bold = True
            paragraph.add_run("Client Ltd in 2025")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "OldClient budget"
            document.sections[0].header.paragraphs[0].text = "OldClient confidential"
            document.save(root / "proposal.docx")
            output, count = replace_text(root / "proposal.docx", root / "edited.docx",
                                         {"OldClient": "NewClient", "2025": "2026"})
            self.assertEqual(count, 4)
            edited = Document(output)
            self.assertEqual(edited.paragraphs[0].text, "Prepared for NewClient Ltd in 2026")
            self.assertTrue(edited.paragraphs[0].runs[0].bold)
            self.assertEqual(edited.tables[0].cell(0, 0).text, "NewClient budget")
            self.assertEqual(edited.sections[0].header.paragraphs[0].text, "NewClient confidential")
            self.assertIn("OldClient", Document(root / "proposal.docx").paragraphs[0].text)

    def test_pptx_and_xlsx_replace_and_options(self):
        from openpyxl import Workbook, load_workbook
        from pptx import Presentation
        from office_edit import replace_text
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = "Acme pitch"
            slide.placeholders[1].text = "acme results and Acmeville"
            deck.save(root / "pitch.pptx")
            output, count = replace_text(root / "pitch.pptx", root / "pitch-edited.pptx", {"acme": "Zenith"},
                                         whole_word=True)
            self.assertEqual(count, 2)
            texts = [s.text_frame.text for s in Presentation(output).slides[0].shapes if s.has_text_frame]
            self.assertEqual(texts, ["Zenith pitch", "Zenith results and Acmeville"])
            book = Workbook()
            book.active.append(["Acme total", "=SUM(1,2)", 5])
            book.save(root / "budget.xlsx")
            book.close()
            output, count = replace_text(root / "budget.xlsx", root / "budget-edited.xlsx", {"Acme": "Zenith"},
                                         match_case=True)
            edited = load_workbook(output)
            try:
                self.assertEqual([c.value for c in edited.active[1]], ["Zenith total", "=SUM(1,2)", 5])
            finally:
                edited.close()
            with self.assertRaises(ValueError):
                replace_text(root / "budget.xlsx", root / "none.xlsx", {"absent": "x"})
            self.assertFalse((root / "none.xlsx").exists())
            with self.assertRaises(ValueError):
                replace_text(root / "budget.xlsx", root / "wrong.docx", {"Acme": "x"})
            with self.assertRaises(ValueError):
                replace_text(root / "budget.xlsx", root / "bad.xlsx", {"(": "x"}, regex=True)

    def test_properties_roundtrip(self):
        from docx import Document
        from openpyxl import Workbook
        from office_edit import read_properties, write_properties
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            Document().save(root / "doc.docx")
            book = Workbook()
            book.save(root / "book.xlsx")
            book.close()
            for name in ("doc.docx", "book.xlsx"):
                output = write_properties(root / name, root / f"tagged-{name}", {"title": "Launch plan", "author": "Studio"})
                properties = read_properties(output)
                self.assertEqual((properties["title"], properties["author"]), ("Launch plan", "Studio"))
            with self.assertRaises(ValueError):
                write_properties(root / "doc.docx", root / "x.docx", {"colour": "red"})
            with self.assertRaises(ValueError):
                read_properties(root / "notes.txt")

    def test_legacy_documents_need_the_office_engine(self):
        from office_edit import replace_text
        with tempfile.TemporaryDirectory() as td, patch("office_engine.find_libreoffice", return_value=None):
            source = Path(td) / "old.doc"
            source.write_bytes(b"legacy")
            with self.assertRaises(RuntimeError) as caught:
                replace_text(source, Path(td) / "new.docx", {"a": "b"})
            self.assertIn("Office engine is missing", str(caught.exception))


class PdfLayoutRouteTests(unittest.TestCase):
    def test_pdf_targets_and_fallback_without_office_engine(self):
        from reportlab.pdfgen import canvas
        from doc_converter import convert_document, document_targets
        self.assertIn("PPTX", document_targets(Path("deck.pdf")))
        self.assertIn("XLSX", document_targets(Path("deck.pdf")))
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            sheet = canvas.Canvas(str(root / "brief.pdf"))
            sheet.drawString(72, 760, "Campaign brief text")
            sheet.save()
            with patch("doc_converter.find_libreoffice", return_value=None), \
                    patch("office_engine.find_libreoffice", return_value=None):
                reflow = convert_document(root / "brief.pdf", root / "out", "DOCX")
                self.assertEqual(reflow.status, "Completed", reflow.error)
                self.assertIn("Text conversion", reflow.note)
                missing = convert_document(root / "brief.pdf", root / "out", "PPTX")
                self.assertEqual(missing.status, "Failed")
                self.assertIn("Office engine is missing", missing.error)
            no_tables = convert_document(root / "brief.pdf", root / "out", "XLSX")
            self.assertEqual(no_tables.status, "Failed")
            self.assertIn("No tables", no_tables.error)
            self.assertEqual([p.name for p in (root / "out").iterdir()], ["brief.docx"])


@unittest.skipUnless(os.environ.get("SHADOW_TEST_FULL_STUDIO") == "1", "Set SHADOW_TEST_FULL_STUDIO=1 for actual local engines")
class PdfLayoutEngineTests(unittest.TestCase):
    def test_pdf_to_docx_and_pptx_keep_pages_and_text(self):
        from docx import Document
        from pptx import Presentation
        from reportlab.pdfgen import canvas
        from doc_converter import convert_document
        from office_edit import replace_text
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            sheet = canvas.Canvas(str(root / "brief.pdf"))
            sheet.drawString(72, 760, "Campaign brief")
            sheet.showPage()
            sheet.drawString(72, 760, "Second page")
            sheet.save()
            word = convert_document(root / "brief.pdf", root / "out", "DOCX")
            self.assertEqual(word.status, "Completed", word.error)
            self.assertIn("layout import", word.note)
            deck = convert_document(root / "brief.pdf", root / "out", "PPTX")
            self.assertEqual(deck.status, "Completed", deck.error)
            self.assertEqual(len(Presentation(deck.output_path).slides), 2)
            # Imported text frames stay editable.
            edited, count = replace_text(word.output_path, root / "edited.docx", {"Campaign brief": "Launch brief"})
            self.assertGreaterEqual(count, 1)
            self.assertTrue(Document(edited))


if __name__ == "__main__":
    unittest.main()
