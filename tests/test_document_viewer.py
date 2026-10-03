from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import _headless
import customtkinter as ctk
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from document_preview import DocumentPreviewDialog
from document_cache import preview_pdf
from ui_dispatch import cancel_widget_callbacks


class DocumentCacheTests(unittest.TestCase):
    def test_reopen_uses_cache_and_source_change_invalidates_it(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);source=root/'report.docx';source.write_bytes(b'first')
            def render(source,destination,*args):destination.write_bytes(b'%PDF fixture')
            with patch('document_cache.get_app_data_dir',return_value=root/'cache'),patch('document_cache.render_office',side_effect=render) as convert:
                self.assertFalse(preview_pdf(source)[1]);self.assertTrue(preview_pdf(source)[1])
                self.assertEqual(convert.call_count,1)
                source.write_bytes(b'changed source')
                self.assertFalse(preview_pdf(source)[1]);self.assertEqual(convert.call_count,2)


class DocumentViewerTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=ctk.CTk();self.root.withdraw();self.addCleanup(self.close)
        self.dialog=None

    def close(self):
        if self.dialog:self.dialog.destroy()
        cancel_widget_callbacks(self.root);self.root.destroy()

    def pump_until(self,condition):
        deadline=time.monotonic()+5
        while not condition() and time.monotonic()<deadline:self.root.update();time.sleep(.01)
        self.assertTrue(condition())

    def test_first_page_does_not_wait_for_text_and_page_back_is_cached(self):
        source=Path(self.tmp.name)/'brief.pdf';source.write_bytes(b'fixture')
        with patch('document_preview.render_pdf_page',return_value=(Image.new('RGB',(600,800),'white'),2)) as render,patch('document_preview.to_markdown',return_value='Extracted brief') as extract:
            self.dialog=DocumentPreviewDialog(self.root,source)
            self.pump_until(lambda:not self.dialog._busy)
            extract.assert_not_called()
            self.dialog._start_text();self.pump_until(lambda:self.dialog._text_loaded)
            self.assertIn('Extracted brief',self.dialog.text.get('1.0','end'))
            self.dialog._navigate(1);self.pump_until(lambda:not self.dialog._busy)
            self.dialog._navigate(-1)
            self.assertEqual(render.call_count,2)

    def test_workbook_sheets_open_without_office_render_and_source_stays_unchanged(self):
        import hashlib
        source=Path(self.tmp.name)/'budget.xlsx';book=Workbook();sheet=book.active;sheet.title='Budget'
        sheet['A1']='Client budget';sheet['A1'].font=Font(bold=True,color='FFFFFF');sheet['A1'].fill=PatternFill('solid',fgColor='7561D4')
        sheet['B2']=1250;sheet['B2'].number_format='$#,##0.00';book.create_sheet('Schedule')['A1']='Launch';book.save(source)
        before=hashlib.sha256(source.read_bytes()).hexdigest()
        with patch('document_preview.preview_pdf') as render:
            self.dialog=DocumentPreviewDialog(self.root,source)
            self.pump_until(lambda:self.dialog.workbook._book is not None)
            self.assertEqual(self.dialog.tabs.get(),'Sheets');render.assert_not_called()
            self.dialog.workbook.show_sheet('Schedule')
            self.assertEqual(self.dialog.workbook._sheet['A1'].value,'Launch')
            self.dialog.workbook.show_sheet('Budget')
            from workbook_view import display_value
            self.assertEqual(display_value(self.dialog.workbook._sheet['B2']),'$1,250.00')
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),before)
