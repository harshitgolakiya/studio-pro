from __future__ import annotations

import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import _headless

from doc_converter import convert_document, document_targets, to_markdown
from office_engine import find_libreoffice, render_office
from pdf_tools import parse_pages, process_pdf, render_pdf_page
from speech_engine import TranscriptSegment, format_transcript, synthesize_speech, transcribe_media
from studio_runtime import run_engine, StudioCancelled


class DocumentRoutesTests(unittest.TestCase):
    def test_targets_follow_document_family(self):
        self.assertIn("PPTX", document_targets(Path("pitch.ppt")))
        self.assertNotIn("DOCX", document_targets(Path("pitch.ppt")))
        self.assertIn("CSV", document_targets(Path("budget.xlsx")))
        self.assertNotIn("PPTX", document_targets(Path("proposal.doc")))

    def test_invalid_target_never_creates_output(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "proposal.docx"
            source.write_bytes(b"invalid")
            result = convert_document(source, Path(td) / "out", "PPTX")
            self.assertEqual(result.status, "Failed")
            self.assertIn("cannot be converted", result.error)
            self.assertFalse(list((Path(td) / "out").glob("*")))

    def test_unknown_target_is_not_silently_markdown(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "notes.txt"
            source.write_text("hello")
            result = convert_document(source, Path(td), "UNKNOWN")
            self.assertEqual(result.status, "Failed")
            self.assertFalse((Path(td) / "notes.md").exists())

    def test_document_replacement_only_after_success(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "notes.txt"
            source.write_text("hello")
            result = convert_document(source, Path(td), "MD", replace_source=True)
            self.assertEqual(result.status, "Completed")
            self.assertFalse(source.exists())
            self.assertEqual(result.output_path.read_text(), "hello")
            broken = Path(td) / "broken.docx"
            broken.write_bytes(b"broken")
            result = convert_document(broken, Path(td), "MD", replace_source=True)
            self.assertEqual(result.status, "Failed")
            self.assertTrue(broken.exists())

    def test_missing_office_engine_has_actionable_error(self):
        with tempfile.TemporaryDirectory() as td, patch("office_engine.find_libreoffice", return_value=None):
            source = Path(td) / "pitch.ppt"
            source.write_bytes(b"dummy")
            result = convert_document(source, Path(td), "PDF")
            self.assertEqual(result.status, "Failed")
            self.assertIn("Office engine is missing", result.error)
            self.assertFalse((Path(td) / "pitch.pdf").exists())

    def test_presentation_and_workbook_text_extraction(self):
        from pptx import Presentation
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Agency Pitch"
            deck.save(root / "pitch.pptx")
            self.assertIn("Agency Pitch", to_markdown(root / "pitch.pptx"))
            workbook = Workbook()
            workbook.active.append(["Budget", 200])
            workbook.save(root / "budget.xlsx")
            workbook.close()
            self.assertIn("Budget | 200", to_markdown(root / "budget.xlsx"))


class PdfToolsTests(unittest.TestCase):
    def test_ranges_preserve_order_and_reject_bad_pages(self):
        self.assertEqual(parse_pages("3,1-2,3", 4), [2, 0, 1])
        for value in ("0", "9", "3-1", "x", "1-2-3", "1,"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_pages(value, 4)

    def test_merge_extract_rotate_and_render(self):
        from pypdf import PdfWriter, PdfReader
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            sources = []
            for i in range(2):
                source = root / f"input-{i}.pdf"
                with PdfWriter() as writer:
                    writer.add_blank_page(width=300 + i, height=400)
                    writer.write(source)
                sources.append(source)
            merged = process_pdf(sources, root / "merged.pdf")
            self.assertEqual(len(PdfReader(merged).pages), 2)
            extracted = process_pdf([merged], root / "page.pdf", "extract", "2")
            self.assertEqual(len(PdfReader(extracted).pages), 1)
            rotated = process_pdf([merged], root / "rotated.pdf", "rotate", "1")
            self.assertEqual(PdfReader(rotated).pages[0].rotation, 90)
            self.assertEqual(len(PdfReader(rotated).pages), 2)
            self.assertEqual(PdfReader(rotated).pages[1].rotation, 0)
            writer = PdfWriter(clone_from=merged)
            try:
                writer.add_metadata({"/Title": "Agency report"})
                writer.add_outline_item("Campaign", 1)
                writer.write(root / "report.pdf")
            finally:
                writer.close()
            compressed = process_pdf([root / "report.pdf"], root / "compressed.pdf", "compress", "1")
            reader = PdfReader(compressed)
            self.assertEqual(len(reader.pages), 2)
            self.assertEqual(reader.metadata.title, "Agency report")
            self.assertEqual(reader.outline[0].title, "Campaign")
            image, count = render_pdf_page(merged, 1)
            self.assertEqual(count, 2)
            self.assertGreater(image.width, 0)
            image.close()
            archive = process_pdf([merged], root / "pages.zip", "images", "2")
            with zipfile.ZipFile(archive) as z:
                self.assertEqual(z.namelist(), ["page-0002.png"])
            with self.assertRaises(ValueError):
                process_pdf([merged], merged, "extract")

    def test_cancel_keeps_original_and_has_no_output(self):
        with tempfile.TemporaryDirectory() as td:
            destination = Path(td) / "result.pdf"
            with self.assertRaises(StudioCancelled):
                process_pdf([Path(td) / "source.pdf"], destination, cancel_check=lambda: True)
            self.assertFalse(destination.exists())


class SpeechTests(unittest.TestCase):
    def test_subtitle_timestamps_and_formats(self):
        segments = [TranscriptSegment(1.234, 62.345, "Hello agency")]
        self.assertIn("00:00:01,234 --> 00:01:02,345", format_transcript(segments, "SRT"))
        self.assertTrue(format_transcript(segments, "VTT").startswith("WEBVTT"))
        self.assertIn("00:00:01.234", format_transcript(segments, "VTT"))
        self.assertEqual(format_transcript(segments, "TXT"), "Hello agency\n")

    def test_missing_model_fails_without_network(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "recording.wav"
            source.write_bytes(b"audio")
            result = transcribe_media(source, Path(td), model_path=Path(td) / "missing")
            self.assertEqual(result.status, "Failed")
            self.assertIn("model missing", result.error)

    def test_invalid_synthesis_settings_do_not_create_output(self):
        with tempfile.TemporaryDirectory() as td:
            destination = Path(td) / "voice.wav"
            with self.assertRaises(ValueError):
                synthesize_speech("", destination)
            with self.assertRaises(ValueError):
                synthesize_speech("hello", destination, speed=0)
            self.assertFalse(destination.exists())


class RuntimeTests(unittest.TestCase):
    def test_cancellable_process_runner(self):
        self.assertIn("ready", run_engine([sys.executable, "-c", "print('ready')"]))
        with self.assertRaises(TimeoutError):
            run_engine([sys.executable, "-c", "import time; time.sleep(10)"], timeout=0.1)
        with self.assertRaises(StudioCancelled):
            run_engine([sys.executable, "-c", "print('no')"], cancel_check=lambda: True)

    def test_audio_cleanup_and_trim(self):
        import wave
        from audio_tools import process_audio
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "recording.wav"
            with wave.open(str(source), "wb") as wav:
                wav.setnchannels(1)
                wav.setsampwidth(2)
                wav.setframerate(16000)
                wav.writeframes(b"\x00\x00" * 32000)
            output = process_audio(source, root / "clean.wav", denoise=True, start=0.5, duration=1)
            with wave.open(str(output)) as wav:
                self.assertAlmostEqual(wav.getnframes() / wav.getframerate(), 1, places=1)
                self.assertFalse(any(wav.readframes(wav.getnframes())), "Silence must remain silent")
            mp3 = process_audio(source, root / "clean.mp3", "MP3", denoise=True)
            self.assertGreater(mp3.stat().st_size, 100)
            import array
            import math
            tone = root / "tone.wav"
            samples = array.array("h", (int(1000 * math.sin(2 * math.pi * 440 * i / 16000)) for i in range(32000)))
            with wave.open(str(tone), "wb") as wav:
                wav.setnchannels(1)
                wav.setsampwidth(2)
                wav.setframerate(16000)
                wav.writeframes(samples.tobytes())
            normalized = process_audio(tone, root / "tone-normalized.wav")
            with wave.open(str(normalized)) as wav:
                output_samples = array.array("h", wav.readframes(wav.getnframes()))
                peak = max(abs(sample) for sample in output_samples)
                self.assertGreater(peak, 1000)
                self.assertLess(peak, 32767)
            self.assertTrue(source.exists())
            with self.assertRaises(ValueError):
                process_audio(source, source)
            with self.assertRaises(ValueError):
                process_audio(source, root / "bad.wav", start=-1)


@unittest.skipUnless(os.environ.get("SHADOW_TEST_FULL_STUDIO") == "1", "Set SHADOW_TEST_FULL_STUDIO=1 for actual local engines")
class LocalEngineIntegrationTests(unittest.TestCase):
    def test_office_exports_and_legacy_roundtrip(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        from pypdf import PdfReader
        self.assertIsNotNone(find_libreoffice())
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            doc = Document()
            doc.add_heading("Agency proposal", 0)
            doc.add_paragraph("Client campaign scope")
            doc.add_page_break()
            doc.add_paragraph("Second proposal page")
            doc.save(root / "proposal.docx")
            deck = Presentation()
            for title in ("Agency pitch", "Campaign results"):
                deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = title
            deck.save(root / "pitch.pptx")
            book = Workbook()
            book.active.append(["Budget", 2500])
            book.save(root / "budget.xlsx")
            book.close()
            for name, page_count in (("proposal.docx", 2), ("pitch.pptx", 2), ("budget.xlsx", 1)):
                result = convert_document(root / name, root / "out", "PDF")
                self.assertEqual(result.status, "Completed", result.error)
                self.assertEqual(len(PdfReader(result.output_path).pages), page_count)
            for name, legacy in (("proposal.docx", "DOC"), ("pitch.pptx", "PPT"), ("budget.xlsx", "XLS")):
                result = convert_document(root / name, root / "legacy", legacy)
                self.assertEqual(result.status, "Completed", result.error)
                pdf = convert_document(result.output_path, root / "legacy-pdf", "PDF")
                self.assertEqual(pdf.status, "Completed", pdf.error)
            csv = convert_document(root / "budget.xlsx", root / "out", "CSV")
            self.assertEqual(csv.status, "Completed", csv.error)
            self.assertIn("Budget", csv.output_path.read_text(encoding="utf-8"))

    def test_local_voice_to_transcript_roundtrip(self):
        import wave
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            audio = synthesize_speech("Hello. This is the agency studio. We create marketing campaigns and client presentations.", root / "voice.wav")
            with wave.open(str(audio)) as wav:
                self.assertGreater(wav.getnframes(), 1000)
            result = transcribe_media(audio, root, "SRT")
            self.assertEqual(result.status, "Completed", result.error)
            text = result.output_path.read_text(encoding="utf-8").lower()
            self.assertIn("agency", text)
            self.assertIn("-->", text)


class StudioGuiTests(unittest.TestCase):
    def test_document_preview_load_and_callback_cleanup(self):
        import time
        import main
        from document_preview import DocumentPreviewDialog
        with tempfile.TemporaryDirectory() as td, patch("settings.SETTINGS_FILE", Path(td) / "settings.json"):
            app = main.WebPCompressorApp()
            app.withdraw()
            try:
                source = Path(td) / "notes.txt"
                source.write_text("Agency preview text", encoding="utf-8")
                preview = DocumentPreviewDialog(app, source)
                preview.withdraw()
                deadline = time.monotonic() + 10
                while preview._busy and time.monotonic() < deadline:
                    app.update()
                    time.sleep(0.01)
                self.assertFalse(preview._busy)
                self.assertIn("Agency preview text", preview.text.get("1.0", "end"))
                preview.destroy()
                app.update()
            finally:
                app.destroy()

    def test_document_family_switch_and_tools_window(self):
        import main
        with tempfile.TemporaryDirectory() as td, patch("settings.SETTINGS_FILE", Path(td) / "settings.json"):
            app = main.WebPCompressorApp()
            app.withdraw()
            try:
                sources = [Path(td) / "pitch.pptx", Path(td) / "budget.xlsx"]
                for source in sources:
                    source.write_bytes(b"placeholder")
                sources = [source.resolve() for source in sources]
                app._ingest_image_paths(sources)
                app.table.selection_set(app.row_ids[sources[0]])
                app._adapt_settings_to_selection()
                self.assertIn("Document: PPTX", app.format_menu.cget("values"))
                app.table.selection_set(app.row_ids[sources[1]])
                app._adapt_settings_to_selection()
                self.assertIn("Document: XLSX", app.format_menu.cget("values"))
                self.assertNotIn("Document: PPTX", app.format_menu.cget("values"))
                tools = app._open_studio_tools()
                self.assertTrue(tools.winfo_exists())
                tools.destroy()
            finally:
                app.destroy()


if __name__ == "__main__":
    unittest.main()
