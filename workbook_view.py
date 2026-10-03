"""Read-only, virtualized workbook grid with sheet navigation and cell styles."""
from bisect import bisect_right
from datetime import date, datetime
from pathlib import Path
import re
import threading
import tkinter as tk
import tkinter.font as tkfont
import customtkinter as ctk


def display_value(cell):
    value = cell.value
    if value is None:return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d %H:%M' if isinstance(value, datetime) and (value.hour or value.minute) else '%Y-%m-%d')
    if isinstance(value, bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value, (float, int)):
        fmt = cell.number_format
        if fmt == 'General':return f'{value:g}'
        pattern = fmt.split(';')[0]
        decimals = len(re.search(r'\.([0#]+)', pattern).group(1)) if re.search(r'\.([0#]+)', pattern) else 0
        if '%' in pattern:return f'{value*100:.{decimals}f}%'
        if any(symbol in pattern for symbol in ('$', '£', '€')):
            symbol = next(symbol for symbol in ('$', '£', '€') if symbol in pattern)
            return f'{symbol}{value:,.{decimals}f}'
        if ',' in pattern:return f'{value:,.{decimals}f}'
        if '.' in pattern:return f'{value:.{decimals}f}'
        return str(value)
    return str(value)


def rgb(color, fallback):
    if color and color.type == 'rgb' and isinstance(color.rgb, str):return '#' + color.rgb[-6:]
    return fallback


class WorkbookView(ctk.CTkFrame):
    def __init__(self, master, source, cancel, bridge, status):
        super().__init__(master, fg_color='transparent')
        self.source, self.cancel, self.bridge, self.status = Path(source), cancel, bridge, status
        self._book = None;self._sheet = None;self._draw_job = None
        self._fonts = {};self._selection = None
        self.grid_columnconfigure(0, weight=1);self.grid_rowconfigure(1, weight=1)
        self.details = ctk.CTkLabel(self, text='Read-only workbook · loading sheets…', anchor='w', wraplength=650)
        self.details.grid(row=0, column=0, columnspan=2, sticky='ew', pady=(0,6))
        self.canvas = tk.Canvas(self, background='#FFFFFF', highlightthickness=0, xscrollincrement=30, yscrollincrement=22)
        self.canvas.grid(row=1, column=0, sticky='nsew')
        vertical = ctk.CTkScrollbar(self, orientation='vertical', command=lambda *args:self._scroll('y', args))
        vertical.grid(row=1, column=1, sticky='ns')
        horizontal = ctk.CTkScrollbar(self, orientation='horizontal', command=lambda *args:self._scroll('x', args))
        horizontal.grid(row=2, column=0, sticky='ew')
        self.canvas.configure(xscrollcommand=horizontal.set, yscrollcommand=vertical.set)
        self.canvas.bind('<Configure>', self._schedule_draw)
        self.canvas.bind('<MouseWheel>', self._wheel)
        self.canvas.bind('<Button-4>', lambda event:self._linux_wheel(-3))
        self.canvas.bind('<Button-5>', lambda event:self._linux_wheel(3))
        self.canvas.bind('<Button-1>', self._select_cell)
        self.sheet_menu = ctk.CTkOptionMenu(self, values=['Loading…'], command=self.show_sheet, width=220)
        self.sheet_menu.grid(row=3, column=0, sticky='w', pady=(8,0))
        threading.Thread(target=self._load, daemon=True).start()

    def _load(self):
        from openpyxl import load_workbook
        try:
            book = load_workbook(self.source, data_only=True, keep_links=False)
            if self.cancel.is_set():book.close();return
            self.bridge.post(lambda:self._loaded(book))
        except Exception as exc:
            message = str(exc)
            self.bridge.post(lambda:self.status.set(f'Workbook grid unavailable: {message}. Use Pages for its layout.'))

    def _loaded(self, book):
        self._book = book
        names = [sheet.title for sheet in book.worksheets if sheet.sheet_state == 'visible'] or book.sheetnames
        self.sheet_menu.configure(values=names);self.sheet_menu.set(names[0]);self.show_sheet(names[0])

    def show_sheet(self, name):
        if self._book is None:return
        from openpyxl.utils import get_column_letter
        self._sheet = self._book[name];self._selection = None
        scale = self._get_widget_scaling()
        self._row_height = 24*scale;self._header_height = 28*scale;self._header_width = 46*scale
        self._rows = max(30,self._sheet.max_row)
        self._columns = max(12,self._sheet.max_column)
        self._offsets = [0]
        for column in range(1,self._columns+1):
            dimension = self._sheet.column_dimensions.get(get_column_letter(column))
            width = 0 if dimension and dimension.hidden else ((dimension.width if dimension and dimension.width else 12)*7+8)*scale
            self._offsets.append(self._offsets[-1]+width)
        self.canvas.configure(scrollregion=(0,0,self._offsets[-1]+self._header_width,self._rows*self._row_height+self._header_height))
        self.canvas.xview_moveto(0);self.canvas.yview_moveto(0)
        self.details.configure(text=f'{name} · {self._sheet.max_row:,} rows × {self._sheet.max_column:,} columns · Read-only')
        self.status.set('Read-only sheets · formula results use saved workbook values · Pages shows charts and print layout')
        self._schedule_draw()

    def _scroll(self, axis, args):
        (self.canvas.xview if axis=='x' else self.canvas.yview)(*args);self._schedule_draw()

    def _wheel(self, event):
        delta = -int(event.delta/120) if abs(event.delta)>=120 else -1 if event.delta>0 else 1
        (self.canvas.xview_scroll if event.state & 1 else self.canvas.yview_scroll)(delta*3,'units')
        self._schedule_draw();return 'break'

    def _linux_wheel(self, units):
        self.canvas.yview_scroll(units,'units');self._schedule_draw();return 'break'

    def _schedule_draw(self, event=None):
        if self._draw_job is None:self._draw_job=self.after(20,self._render_grid)

    def _font(self, cell=None, header=False):
        scale=self._get_widget_scaling()
        key=('Segoe UI',int(13*scale),True,False) if header else (cell.font.name or 'Calibri',int((cell.font.sz or 11)*1.25*scale),bool(cell.font.bold),bool(cell.font.italic))
        if key not in self._fonts:self._fonts[key]=tkfont.Font(self.canvas,family=key[0],size=-key[1],weight='bold' if key[2] else 'normal',slant='italic' if key[3] else 'roman')
        return self._fonts[key]

    @staticmethod
    def _fit(text, font, width):
        if font.measure(text)<=width:return text
        while text and font.measure(text+'…')>width:text=text[:-1]
        return text+'…'

    def _render_grid(self):
        self._draw_job=None
        if self._sheet is None or self.cancel.is_set():return
        from openpyxl.utils import get_column_letter
        from openpyxl.cell.cell import Cell
        canvas=self.canvas;canvas.delete('all')
        left,top=canvas.canvasx(0),canvas.canvasy(0)
        right,bottom=left+canvas.winfo_width(),top+canvas.winfo_height()
        first_row=max(0,int((top-self._header_height)/self._row_height))
        last_row=min(self._rows,int(bottom/self._row_height)+1)
        first_col=max(0,bisect_right(self._offsets,max(0,left-self._header_width))-1)
        last_col=min(self._columns,bisect_right(self._offsets,right-self._header_width)+1)
        for row in range(first_row,last_row):
            y=self._header_height+row*self._row_height
            for column in range(first_col,last_col):
                x=self._header_width+self._offsets[column];width=self._offsets[column+1]-self._offsets[column]
                if width<1:continue
                cell=self._sheet._cells.get((row+1,column+1)) or Cell(self._sheet,row=row+1,column=column+1)
                fill=rgb(cell.fill.fgColor,'#FFFFFF') if cell.fill.patternType=='solid' else '#FFFFFF'
                canvas.create_rectangle(x,y,x+width,y+self._row_height,fill=fill,outline='#E3E5E9')
                text=display_value(cell)
                if text:
                    font=self._font(cell);text=self._fit(text,font,max(0,width-10))
                    numeric=isinstance(cell.value,(int,float,date,datetime))
                    alignment=cell.alignment.horizontal
                    anchor='e' if alignment=='right' or (numeric and not alignment) else 'center' if alignment=='center' else 'w'
                    tx=x+width-5 if anchor=='e' else x+width/2 if anchor=='center' else x+5
                    canvas.create_text(tx,y+self._row_height/2,text=text,anchor=anchor,font=font,fill=rgb(cell.font.color,'#25252F'))
                if self._selection==(row,column):canvas.create_rectangle(x,y,x+width,y+self._row_height,outline='#7561D4',width=2)
        for merged in self._sheet.merged_cells.ranges:
            if merged.max_row<first_row+1 or merged.min_row>last_row or merged.max_col<first_col+1 or merged.min_col>last_col:continue
            cell=self._sheet.cell(merged.min_row,merged.min_col)
            x=self._header_width+self._offsets[merged.min_col-1];x2=self._header_width+self._offsets[merged.max_col]
            y=self._header_height+(merged.min_row-1)*self._row_height;y2=self._header_height+merged.max_row*self._row_height
            fill=rgb(cell.fill.fgColor,'#FFFFFF') if cell.fill.patternType=='solid' else '#FFFFFF'
            canvas.create_rectangle(x,y,x2,y2,fill=fill,outline='#E3E5E9')
            font=self._font(cell);text=self._fit(display_value(cell),font,max(0,x2-x-10))
            center=cell.alignment.horizontal=='center'
            canvas.create_text((x+x2)/2 if center else x+5,(y+y2)/2,text=text,font=font,anchor='center' if center else 'w',fill=rgb(cell.font.color,'#25252F'))
        for row in range(first_row,last_row):
            y=self._header_height+row*self._row_height
            canvas.create_rectangle(left,y,left+self._header_width,y+self._row_height,fill='#F0F1F5',outline='#D9DBE3')
            canvas.create_text(left+self._header_width/2,y+self._row_height/2,text=str(row+1),font=self._font(header=True),fill='#555565')
        canvas.create_rectangle(left,top,right,top+self._header_height,fill='#F0F1F5',outline='#D9DBE3')
        for column in range(first_col,last_col):
            x=self._header_width+self._offsets[column];width=self._offsets[column+1]-self._offsets[column]
            if width<1:continue
            canvas.create_rectangle(x,top,x+width,top+self._header_height,fill='#F0F1F5',outline='#D9DBE3')
            canvas.create_text(x+width/2,top+self._header_height/2,text=get_column_letter(column+1),font=self._font(header=True),fill='#555565')
        canvas.create_rectangle(left,top,left+self._header_width,top+self._header_height,fill='#E6E7ED',outline='#D9DBE3')

    def _select_cell(self,event):
        if self._sheet is None or event.x<self._header_width or event.y<self._header_height:return
        row=int((self.canvas.canvasy(event.y)-self._header_height)/self._row_height)
        column=bisect_right(self._offsets,self.canvas.canvasx(event.x)-self._header_width)-1
        if 0<=row<self._rows and 0<=column<self._columns:
            self._selection=(row,column);cell=self._sheet.cell(row+1,column+1)
            self.details.configure(text=f'{cell.coordinate} · {display_value(cell)}');self._schedule_draw()

    def destroy(self):
        if self._draw_job:
            try:self.after_cancel(self._draw_job)
            except tk.TclError:pass
        if self._book:self._book.close()
        super().destroy()
