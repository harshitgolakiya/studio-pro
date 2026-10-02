"""PDF passwords, redaction, forms, signatures, repair, and table extraction.

Every operation writes a new file and leaves the source untouched.
"""
from __future__ import annotations

import csv
import ctypes
import io
import json
from pathlib import Path
import re
from typing import Callable

from pdf_tools import _render_lock, parse_pages, read_pdf
from studio_runtime import StagedOutput, check_cancel

REDACTION_PRESETS = {
    "Email addresses": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
    "Phone numbers": r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)",
    "Card-like numbers": r"(?<!\d)(?:\d[ -]?){13,19}(?!\d)",
}
TABLE_FORMATS = {"XLSX": ".xlsx", "CSV": ".csv", "JSON": ".json"}


def _require_pdf(source: Path, destination: Path, extension: str = ".pdf") -> None:
    if not source.is_file():
        raise FileNotFoundError("The selected PDF no longer exists")
    if destination.suffix.lower() != extension:
        raise ValueError(f"This operation requires a {extension} output")
    if destination.resolve() == source.resolve():
        raise ValueError("Choose a new output path to keep your original PDF")


# -- passwords -----------------------------------------------------------

def unlock_pdf(source: Path, destination: Path, password: str, overwrite: bool = False) -> Path:
    """Save an unencrypted copy of a PDF the user holds the password for."""
    from pypdf import PdfWriter
    _require_pdf(source, destination)
    reader = read_pdf(source, password)
    if not reader.is_encrypted:
        raise ValueError("This PDF is not password-protected")
    with StagedOutput(destination, overwrite) as stage:
        writer = PdfWriter(clone_from=reader)
        try:
            with stage.path.open("wb") as stream:
                writer.write(stream)
        finally:
            writer.close()
    return stage.output


def protect_pdf(source: Path, destination: Path, open_password: str, owner_password: str = "",
                allow_printing: bool = True, allow_copying: bool = False, allow_editing: bool = False,
                current_password: str | None = None, overwrite: bool = False) -> Path:
    """Encrypt a PDF with AES-256. Viewers must enter ``open_password``."""
    from pypdf import PdfWriter
    from pypdf.constants import UserAccessPermissions as Permission
    _require_pdf(source, destination)
    if not open_password and not owner_password:
        raise ValueError("Enter a password")
    permissions = Permission(0)
    if allow_printing:
        permissions |= Permission.PRINT | Permission.PRINT_TO_REPRESENTATION
    if allow_copying:
        permissions |= Permission.EXTRACT | Permission.EXTRACT_TEXT_AND_GRAPHICS
    if allow_editing:
        permissions |= (Permission.MODIFY | Permission.ADD_OR_MODIFY | Permission.FILL_FORM_FIELDS
                        | Permission.ASSEMBLE_DOC)
    reader = read_pdf(source, current_password)
    with StagedOutput(destination, overwrite) as stage:
        writer = PdfWriter(clone_from=reader)
        try:
            writer.encrypt(open_password, owner_password or open_password, algorithm="AES-256",
                           permissions_flag=permissions)
            with stage.path.open("wb") as stream:
                writer.write(stream)
        finally:
            writer.close()
    return stage.output


# -- redaction -----------------------------------------------------------

def _patterns(terms: list[str], regex: bool, match_case: bool) -> list[re.Pattern]:
    compiled = []
    for term in terms:
        if not term.strip():
            continue
        try:
            compiled.append(re.compile(term if regex else re.escape(term.strip()), 0 if match_case else re.I))
        except re.error as exc:
            raise ValueError(f"Invalid pattern {term!r}: {exc}")
    return compiled


def redact_pdf(source: Path, destination: Path, terms: list[str] | None = None, regex: bool = False,
               match_case: bool = False, regions: list[tuple[int, float, float, float, float]] | None = None,
               pages: str = "all", password: str | None = None, dpi: int = 200, overwrite: bool = False,
               cancel_check: Callable[[], bool] | None = None,
               progress: Callable[[str], None] | None = None) -> tuple[Path, int]:
    """Black out matching text and regions. Returns the output and the number of areas removed.

    Each page with a redaction is replaced by a flattened image of itself, so
    the covered text, its hidden layers, and its annotations no longer exist in
    the file. Pages without a match are kept as they are. ``regions`` are
    ``(page_number, left, top, right, bottom)`` in points from the top-left of
    the page as displayed.
    """
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    from PIL import ImageDraw
    from pypdf import PdfReader, PdfWriter
    _require_pdf(source, destination)
    patterns = _patterns(terms or [], regex, match_case)
    regions = regions or []
    if not patterns and not regions:
        raise ValueError("Enter text to redact or choose a preset")
    if not 100 <= dpi <= 400:
        raise ValueError("Redaction resolution must be between 100 and 400 DPI")
    reader = read_pdf(source, password)
    selected = set(parse_pages(pages, len(reader.pages)))
    replaced: dict[int, bytes] = {}
    total = 0
    with _render_lock:
        with pdfium.PdfDocument(str(source), password=password) as document:
            for index in sorted(selected):
                check_cancel(cancel_check)
                if progress:
                    progress(f"Scanning page {index + 1} of {len(reader.pages)}…")
                page = document[index]
                try:
                    width, height = page.get_size()
                    scale = min(dpi / 72, 4000 / max(width, height))
                    pixel_width, pixel_height = max(1, round(width * scale)), max(1, round(height * scale))
                    boxes = []
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_range()
                        for pattern in patterns:
                            for match in pattern.finditer(text):
                                if match.end() == match.start():
                                    continue
                                for number in range(textpage.count_rects(match.start(), match.end() - match.start())):
                                    left, bottom, right, top = textpage.get_rect(number)
                                    corners = []
                                    for x, y in ((left, top), (right, bottom)):
                                        device_x, device_y = ctypes.c_int(), ctypes.c_int()
                                        raw.FPDF_PageToDevice(page.raw, 0, 0, pixel_width, pixel_height, 0,
                                                              x, y, ctypes.byref(device_x), ctypes.byref(device_y))
                                        corners.append((device_x.value, device_y.value))
                                    boxes.append((min(corners[0][0], corners[1][0]), min(corners[0][1], corners[1][1]),
                                                  max(corners[0][0], corners[1][0]), max(corners[0][1], corners[1][1])))
                    finally:
                        textpage.close()
                    for number, left, top, right, bottom in regions:
                        if number == index + 1:
                            boxes.append((left * scale, top * scale, right * scale, bottom * scale))
                    if not boxes:
                        continue
                    bitmap = page.render(scale=scale)
                    try:
                        image = bitmap.to_pil().convert("RGB")
                    finally:
                        bitmap.close()
                    draw = ImageDraw.Draw(image)
                    pad = max(2, round(scale * 1.5))
                    for left, top, right, bottom in boxes:
                        draw.rectangle((left - pad, top - pad, right + pad, bottom + pad), fill="black")
                    buffer = io.BytesIO()
                    image.save(buffer, format="PDF", resolution=72 * scale, quality=90)
                    image.close()
                    replaced[index] = buffer.getvalue()
                    total += len(boxes)
                finally:
                    page.close()
    if not replaced:
        raise ValueError("Nothing to redact: the text was not found on the selected pages. "
                         "Scanned pages need OCR first (Studio Tools → OCR).")
    with StagedOutput(destination, overwrite) as stage:
        writer = PdfWriter()
        try:
            for index, page in enumerate(reader.pages):
                check_cancel(cancel_check)
                writer.add_page(PdfReader(io.BytesIO(replaced[index])).pages[0] if index in replaced else page)
            # Titles, keywords, and XMP packets can repeat the removed text.
            writer.add_metadata({"/Producer": "Shadow"})
            with stage.path.open("wb") as stream:
                writer.write(stream)
        finally:
            writer.close()
        # Prove the result before publishing it.
        check = PdfReader(stage.path)
        for index in replaced:
            remaining = check.pages[index].extract_text() or ""
            if any(pattern.search(remaining) for pattern in patterns):
                raise RuntimeError(f"Redaction could not be verified on page {index + 1}; no file was written")
    return stage.output, total


# -- forms ---------------------------------------------------------------

def list_form_fields(source: Path, password: str | None = None) -> list[dict]:
    """Fillable fields: name, kind, current value, and the choices for lists and buttons."""
    kinds = {"/Tx": "text", "/Btn": "button", "/Ch": "choice", "/Sig": "signature"}
    fields = []
    for name, field in (read_pdf(source, password).get_fields() or {}).items():
        value = field.get("/V")
        options = field.get("/_States_") or field.get("/Opt") or []
        fields.append({"name": name, "kind": kinds.get(field.get("/FT"), "other"),
                       "value": "" if value is None else str(value),
                       "options": [str(o[1] if isinstance(o, list) else o) for o in options]})
    return fields


def fill_form(source: Path, destination: Path, values: dict[str, str], flatten: bool = False,
              password: str | None = None, overwrite: bool = False) -> Path:
    """Fill form fields by name; ``flatten`` bakes the values into the page."""
    from pypdf import PdfWriter
    _require_pdf(source, destination)
    reader = read_pdf(source, password)
    known = set(reader.get_fields() or {})
    if not known:
        raise ValueError("This PDF has no fillable form fields")
    unknown = sorted(set(values) - known)
    if unknown:
        raise ValueError(f"Unknown form fields: {', '.join(unknown)}. Available: {', '.join(sorted(known))}")
    with StagedOutput(destination, overwrite) as stage:
        writer = PdfWriter(clone_from=reader)
        try:
            writer.update_page_form_field_values(None, {k: str(v) for k, v in values.items()},
                                                 auto_regenerate=False, flatten=flatten)
            if flatten:
                writer.remove_annotations(subtypes="/Widget")
            else:
                writer.set_need_appearances_writer(True)
            with stage.path.open("wb") as stream:
                writer.write(stream)
        finally:
            writer.close()
    return stage.output


# -- signatures ----------------------------------------------------------

def create_signing_identity(name: str, destination: Path, passphrase: str, organization: str = "",
                            email: str = "", years: int = 3, overwrite: bool = False) -> Path:
    """Create a self-signed PKCS#12 identity for signing PDFs.

    Readers show a self-signed signature as "identity not verified" until the
    recipient trusts the certificate; use a CA-issued certificate for public delivery.
    """
    from datetime import datetime, timedelta, timezone
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import ExtendedKeyUsageOID, NameOID
    if not name.strip():
        raise ValueError("Enter the signer's name")
    if len(passphrase) < 6:
        raise ValueError("Use a passphrase of at least 6 characters")
    if destination.suffix.lower() not in {".p12", ".pfx"}:
        raise ValueError("Save the identity as a .p12 or .pfx file")
    key = rsa.generate_private_key(public_exponent=65537, key_size=3072)
    attributes = [x509.NameAttribute(NameOID.COMMON_NAME, name.strip())]
    if organization.strip():
        attributes.append(x509.NameAttribute(NameOID.ORGANIZATION_NAME, organization.strip()))
    if email.strip():
        attributes.append(x509.NameAttribute(NameOID.EMAIL_ADDRESS, email.strip()))
    subject = x509.Name(attributes)
    now = datetime.now(timezone.utc)
    certificate = (x509.CertificateBuilder().subject_name(subject).issuer_name(subject)
                   .public_key(key.public_key()).serial_number(x509.random_serial_number())
                   .not_valid_before(now - timedelta(minutes=5)).not_valid_after(now + timedelta(days=365 * years))
                   .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True)
                   .add_extension(x509.KeyUsage(digital_signature=True, content_commitment=True, key_encipherment=False,
                                                data_encipherment=False, key_agreement=False, key_cert_sign=False,
                                                crl_sign=False, encipher_only=False, decipher_only=False), critical=True)
                   .add_extension(x509.ExtendedKeyUsage([ExtendedKeyUsageOID.EMAIL_PROTECTION]), critical=False)
                   .sign(key, hashes.SHA256()))
    data = pkcs12.serialize_key_and_certificates(
        name.strip().encode("utf-8"), key, certificate, None,
        serialization.BestAvailableEncryption(passphrase.encode("utf-8")))
    with StagedOutput(destination, overwrite) as stage:
        stage.path.write_bytes(data)
    return stage.output


def sign_pdf(source: Path, destination: Path, identity: Path, passphrase: str, reason: str = "",
             location: str = "", field_name: str = "Signature1", overwrite: bool = False) -> Path:
    """Apply a digital signature from a PKCS#12 (.p12/.pfx) identity."""
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    from pyhanko.sign import signers
    _require_pdf(source, destination)
    if not identity.is_file():
        raise FileNotFoundError("Choose a .p12 or .pfx signing identity")
    try:
        signer = signers.SimpleSigner.load_pkcs12(str(identity), passphrase=passphrase.encode("utf-8"))
    except Exception as exc:
        raise ValueError(f"The signing identity could not be opened: {exc}")
    if signer is None:
        raise ValueError("The signing identity could not be opened. Check the passphrase.")
    metadata = signers.PdfSignatureMetadata(field_name=field_name, reason=reason or None, location=location or None)
    with StagedOutput(destination, overwrite) as stage:
        with source.open("rb") as stream, stage.path.open("wb") as output:
            # Signed files are updated incrementally so earlier signatures stay valid.
            signers.sign_pdf(IncrementalPdfFileWriter(stream, strict=False), metadata, signer=signer, output=output)
    return stage.output


def verify_signatures(source: Path) -> list[dict]:
    """Report each signature: signer, whether the signed bytes are intact, and its coverage."""
    from pyhanko.pdf_utils.reader import PdfFileReader
    from pyhanko.sign.validation import validate_pdf_signature
    from pyhanko_certvalidator import ValidationContext
    reports = []
    with source.open("rb") as stream:
        reader = PdfFileReader(stream, strict=False)
        for signature in reader.embedded_signatures:
            report = {"field": signature.field_name, "signer": "", "intact": False, "valid": False,
                      "trusted": False, "covers_whole_document": False, "signed_at": "", "error": ""}
            try:
                # Offline: no certificate revocation or timestamp lookups.
                status = validate_pdf_signature(signature, ValidationContext(allow_fetching=False))
                report.update(signer=status.signing_cert.subject.human_friendly, intact=bool(status.intact),
                              valid=bool(status.valid), trusted=bool(status.trusted),
                              covers_whole_document=status.coverage.name == "ENTIRE_FILE",
                              signed_at=status.signer_reported_dt.isoformat() if status.signer_reported_dt else "")
            except Exception as exc:
                report["error"] = str(exc)
            reports.append(report)
    return reports


# -- repair --------------------------------------------------------------

def repair_pdf(source: Path, destination: Path, password: str | None = None, overwrite: bool = False) -> tuple[Path, int]:
    """Rebuild a damaged PDF's structure. Returns the output and its recovered page count."""
    import pypdfium2 as pdfium
    from pypdf import PdfReader, PdfWriter
    _require_pdf(source, destination)
    with StagedOutput(destination, overwrite) as stage:
        pages = 0
        try:
            # PDFium reconstructs broken cross-reference tables while loading.
            with _render_lock:
                with pdfium.PdfDocument(str(source), password=password) as document:
                    pages = len(document)
                    if pages:
                        document.save(str(stage.path))
        except Exception:
            pages = 0
        if not pages:
            try:
                reader = PdfReader(source, strict=False)
                if reader.is_encrypted and not reader.decrypt(password or ""):
                    raise ValueError("This PDF is password-protected. Enter its password.")
                writer = PdfWriter()
                try:
                    for page in reader.pages:
                        try:
                            writer.add_page(page)
                            pages += 1
                        except Exception:
                            continue
                    if pages:
                        with stage.path.open("wb") as stream:
                            writer.write(stream)
                finally:
                    writer.close()
            except ValueError:
                raise
            except Exception as exc:
                raise RuntimeError(f"This file could not be repaired: {exc}")
        if not pages:
            raise RuntimeError("No pages could be recovered from this file")
        pages = len(PdfReader(stage.path).pages)
    return stage.output, pages


# -- tables --------------------------------------------------------------

def extract_tables(source: Path, destination: Path, pages: str = "all", password: str | None = None,
                   overwrite: bool = False, cancel_check: Callable[[], bool] | None = None,
                   progress: Callable[[str], None] | None = None) -> tuple[Path, int]:
    """Export ruled and aligned tables to XLSX (one sheet each), CSV, or JSON. Returns output and table count."""
    import pdfplumber
    fmt = destination.suffix.lower().lstrip(".").upper()
    if fmt not in TABLE_FORMATS:
        raise ValueError("Save tables as .xlsx, .csv, or .json")
    _require_pdf(source, destination, TABLE_FORMATS[fmt])
    tables: list[dict] = []
    try:
        document = pdfplumber.open(source, password=password or None)
    except Exception as exc:
        raise ValueError(f"This PDF could not be opened: {exc}")
    with document:
        indices = parse_pages(pages, len(document.pages))
        for position, index in enumerate(indices, 1):
            check_cancel(cancel_check)
            if progress:
                progress(f"Reading tables on page {position} of {len(indices)}…")
            for rows in document.pages[index].extract_tables():
                cleaned = [["" if cell is None else str(cell).strip() for cell in row] for row in rows]
                if any(any(row) for row in cleaned):
                    tables.append({"page": index + 1, "table": len(tables) + 1, "rows": cleaned})
    if not tables:
        raise ValueError("No tables were found. Scanned pages need OCR first; "
                         "tables without ruling lines or aligned columns may not be detected.")
    with StagedOutput(destination, overwrite) as stage:
        if fmt == "JSON":
            stage.path.write_text(json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8")
        elif fmt == "CSV":
            with stage.path.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream)
                for number, table in enumerate(tables):
                    if number:
                        writer.writerow([])
                    writer.writerows(table["rows"])
        else:
            from openpyxl import Workbook
            book = Workbook()
            book.remove(book.active)
            for table in tables:
                sheet = book.create_sheet(f"P{table['page']} T{table['table']}"[:31])
                for row in table["rows"]:
                    sheet.append(row)
            book.save(stage.path)
            book.close()
    return stage.output, len(tables)
