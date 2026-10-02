"""Tabular and structured data: CSV, TSV, JSON, JSON Lines, XML, YAML, and workbooks.

Conversion, cleanup, schema validation, and one-file-per-sheet export. All
parsing is local and uses safe loaders (no YAML object construction, no XML
entity expansion).
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
import re
from typing import Any, Callable

from studio_runtime import StagedOutput, check_cancel

DATA_READ_EXTENSIONS = {".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".xml", ".yaml", ".yml", ".xlsx", ".xlsm"}
DATA_WRITE_FORMATS = {"CSV": ".csv", "TSV": ".tsv", "JSON": ".json", "JSONL": ".jsonl",
                      "XML": ".xml", "YAML": ".yaml", "XLSX": ".xlsx"}
_MAX_BYTES = 512 * 1024 * 1024
_XML_NAME = re.compile(r"[^A-Za-z0-9_.-]")


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _typed(value: str) -> Any:
    """Numbers, booleans, and blanks from delimited text; everything else stays text."""
    text = value.strip()
    if text == "":
        return None
    if text.lower() in ("true", "false"):
        return text.lower() == "true"
    # Leading zeros and plus signs mark identifiers (postcodes, phone numbers), not numbers.
    if re.fullmatch(r"-?(0|[1-9]\d*)", text):
        return int(text)
    if re.fullmatch(r"-?(0|[1-9]\d*)\.\d+([eE][-+]?\d+)?", text):
        return float(text)
    return value


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    """Nested objects become dotted columns; lists of scalars join with '; '."""
    if isinstance(value, dict):
        flat: dict[str, Any] = {}
        for key, item in value.items():
            flat.update(_flatten(item, f"{prefix}.{key}" if prefix else str(key)))
        return flat
    if isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            return {prefix: "; ".join("" if item is None else str(item) for item in value)}
        return {prefix: json.dumps(value, ensure_ascii=False)}
    return {prefix: value}


def _records_from(value: Any) -> list[dict[str, Any]]:
    """Find the list of records inside parsed JSON/YAML."""
    if isinstance(value, list):
        return [item if isinstance(item, dict) else {"value": item} for item in value]
    if isinstance(value, dict):
        lists = [item for item in value.values() if isinstance(item, list) and item and isinstance(item[0], dict)]
        if len(lists) == 1:
            return lists[0]
        return [value]
    raise ValueError("This file does not contain records")


def _xml_value(element) -> Any:
    children = list(element)
    if not children and not element.attrib:
        return (element.text or "").strip() or None
    value: dict[str, Any] = {f"@{key}": item for key, item in element.attrib.items()}
    for child in children:
        tag = child.tag.split("}")[-1]
        item = _xml_value(child)
        if tag in value:
            if not isinstance(value[tag], list):
                value[tag] = [value[tag]]
            value[tag].append(item)
        else:
            value[tag] = item
    text = (element.text or "").strip()
    if text:
        value["#text"] = text
    return value


def _sheet_records(sheet) -> list[dict[str, Any]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and all(cell is None for cell in rows[-1]):
        rows.pop()
    if not rows:
        return []
    headers = _unique_headers(["" if cell is None else str(cell) for cell in rows[0]])
    return [{header: (row[i] if i < len(row) else None) for i, header in enumerate(headers)} for row in rows[1:]]


def _unique_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for position, header in enumerate(headers, 1):
        name = header.strip() or f"column_{position}"
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return result


def sheet_names(source: Path) -> list[str]:
    from openpyxl import load_workbook
    book = load_workbook(source, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def load_records(source: Path, sheet: str | None = None, infer_types: bool = True) -> list[dict[str, Any]]:
    """Read any supported data file as a list of records."""
    extension = source.suffix.lower()
    if extension not in DATA_READ_EXTENSIONS:
        raise ValueError(f"Unsupported data file: {extension or source.name}. "
                         f"Supported: {', '.join(sorted(DATA_READ_EXTENSIONS))}")
    if not source.is_file():
        raise FileNotFoundError("The selected data file no longer exists")
    if source.stat().st_size > _MAX_BYTES:
        raise ValueError("This file is larger than 512 MB; split it before converting")
    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError(f"No sheet named {sheet!r}. Available: {', '.join(book.sheetnames)}")
            return _sheet_records(book[sheet] if sheet else book.worksheets[0])
        finally:
            book.close()
    text = _read_text(source)
    try:
        if extension in {".csv", ".tsv"}:
            delimiter = "\t" if extension == ".tsv" else ","
            if extension == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
                except csv.Error:
                    pass
            rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
            if not rows:
                return []
            headers = _unique_headers(rows[0])
            convert = _typed if infer_types else (lambda value: value)
            return [{header: convert(row[i]) if i < len(row) else None for i, header in enumerate(headers)}
                    for row in rows[1:] if row]
        if extension in {".jsonl", ".ndjson"}:
            return _records_from([json.loads(line) for line in text.splitlines() if line.strip()])
        if extension == ".json":
            return _records_from(json.loads(text))
        if extension in {".yaml", ".yml"}:
            import yaml
            return _records_from(yaml.safe_load(text))
        from defusedxml import ElementTree
        root = ElementTree.fromstring(text)
        items = [_xml_value(child) for child in root]
        return [item if isinstance(item, dict) else {"value": item} for item in items] or [_xml_value(root)]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"{source.name} could not be read as {extension.lstrip('.').upper()}: {exc}")


def _columns(records: list[dict[str, Any]]) -> list[str]:
    columns: dict[str, None] = {}
    for record in records:
        for key in record:
            columns.setdefault(key)
    return list(columns)


def _xml_element(parent, name: str, value: Any) -> None:
    from xml.etree import ElementTree
    tag = _XML_NAME.sub("_", name) or "field"
    if tag[0].isdigit() or tag[0] in ".-":
        tag = "_" + tag
    if isinstance(value, list):
        for item in value:
            _xml_element(parent, name, item)
        return
    element = ElementTree.SubElement(parent, tag)
    if isinstance(value, dict):
        for key, item in value.items():
            if key.startswith("@"):
                element.set(_XML_NAME.sub("_", key[1:]) or "attr", "" if item is None else str(item))
            elif key == "#text":
                element.text = str(item)
            else:
                _xml_element(element, key, item)
    elif value is not None:
        element.text = str(value).lower() if isinstance(value, bool) else str(value)


def save_records(records: list[dict[str, Any]], path: Path, fmt: str, sheet_title: str = "Data") -> None:
    """Write records in ``fmt`` to ``path`` (no atomic publish; callers stage the file)."""
    fmt = fmt.upper()
    if fmt not in DATA_WRITE_FORMATS:
        raise ValueError(f"Choose one of: {', '.join(DATA_WRITE_FORMATS)}")
    if fmt in {"CSV", "TSV", "XLSX"}:
        flat = [_flatten(record) for record in records]
        columns = _columns(flat)
        if fmt == "XLSX":
            from openpyxl import Workbook
            book = Workbook(write_only=True)
            sheet = book.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", sheet_title)[:31] or "Data")
            sheet.append(columns)
            for record in flat:
                sheet.append([record.get(column) for column in columns])
            book.save(path)
            book.close()
            return
        with path.open("w", newline="", encoding="utf-8-sig" if fmt == "CSV" else "utf-8") as stream:
            writer = csv.writer(stream, delimiter="\t" if fmt == "TSV" else ",")
            writer.writerow(columns)
            for record in flat:
                writer.writerow(["" if record.get(c) is None else record.get(c) for c in columns])
    elif fmt == "JSON":
        path.write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    elif fmt == "JSONL":
        path.write_text("".join(json.dumps(r, ensure_ascii=False, default=str) + "\n" for r in records), encoding="utf-8")
    elif fmt == "YAML":
        import yaml
        plain = json.loads(json.dumps(records, default=str))
        path.write_text(yaml.safe_dump(plain, allow_unicode=True, sort_keys=False), encoding="utf-8")
    else:
        from xml.etree import ElementTree
        root = ElementTree.Element("records")
        for record in records:
            _xml_element(root, "record", record)
        ElementTree.indent(root)
        ElementTree.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def _format_for(destination: Path) -> str:
    for fmt, extension in DATA_WRITE_FORMATS.items():
        if destination.suffix.lower() == extension or (fmt == "YAML" and destination.suffix.lower() == ".yml"):
            return fmt
    raise ValueError(f"Save data as one of: {', '.join(DATA_WRITE_FORMATS.values())}")


def convert_data(source: Path, destination: Path, sheet: str | None = None, infer_types: bool = True,
                 overwrite: bool = False) -> tuple[Path, int]:
    """Convert between data formats. Returns the output and its record count."""
    fmt = _format_for(destination)
    if destination.resolve() == source.resolve():
        raise ValueError("Choose a new output path to keep your original file")
    records = load_records(source, sheet, infer_types)
    with StagedOutput(destination, overwrite) as stage:
        save_records(records, stage.path, fmt, sheet or source.stem)
    return stage.output, len(records)


# -- cleanup -------------------------------------------------------------

def clean_records(records: list[dict[str, Any]], trim: bool = True, drop_empty_rows: bool = True,
                  drop_empty_columns: bool = True, drop_duplicates: bool = False,
                  normalize_headers: bool = False) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Return cleaned records and a count of each change made."""
    report = {"cells_trimmed": 0, "empty_rows_removed": 0, "empty_columns_removed": 0,
              "duplicate_rows_removed": 0, "headers_renamed": 0}

    def blank(value: Any) -> bool:
        return value is None or (isinstance(value, str) and not value.strip())

    cleaned = []
    for record in records:
        row = dict(record)
        if trim:
            for key, value in row.items():
                if isinstance(value, str):
                    tidy = re.sub(r"[ \t ]+", " ", value).strip()
                    if tidy != value:
                        report["cells_trimmed"] += 1
                        row[key] = tidy
        if drop_empty_rows and all(blank(value) for value in row.values()):
            report["empty_rows_removed"] += 1
            continue
        cleaned.append(row)
    if drop_empty_columns:
        empty = [column for column in _columns(cleaned) if all(blank(row.get(column)) for row in cleaned)]
        report["empty_columns_removed"] = len(empty)
        cleaned = [{key: value for key, value in row.items() if key not in empty} for row in cleaned]
    if drop_duplicates:
        seen, unique = set(), []
        for row in cleaned:
            key = json.dumps(row, sort_keys=True, default=str)
            if key in seen:
                report["duplicate_rows_removed"] += 1
                continue
            seen.add(key)
            unique.append(row)
        cleaned = unique
    if normalize_headers:
        columns = _columns(cleaned)
        renamed = _unique_headers([re.sub(r"[^a-z0-9]+", "_", column.strip().lower()).strip("_") for column in columns])
        mapping = dict(zip(columns, renamed))
        report["headers_renamed"] = sum(1 for old, new in mapping.items() if old != new)
        cleaned = [{mapping[key]: value for key, value in row.items()} for row in cleaned]
    return cleaned, report


def clean_data(source: Path, destination: Path, sheet: str | None = None, overwrite: bool = False,
               **options: bool) -> tuple[Path, dict[str, int]]:
    """Clean a data file into a new file. Returns the output and the change report."""
    fmt = _format_for(destination)
    if destination.resolve() == source.resolve():
        raise ValueError("Choose a new output path to keep your original file")
    # Keep identifiers such as postcodes exactly as typed.
    cleaned, report = clean_records(load_records(source, sheet, infer_types=False), **options)
    report["rows_written"] = len(cleaned)
    with StagedOutput(destination, overwrite) as stage:
        save_records(cleaned, stage.path, fmt, sheet or source.stem)
    return stage.output, report


# -- schema --------------------------------------------------------------

def infer_schema(records: list[dict[str, Any]]) -> dict[str, Any]:
    """A starting JSON Schema describing the records: types seen and always-present fields."""
    names = {bool: "boolean", int: "integer", float: "number", str: "string", list: "array", dict: "object"}
    properties: dict[str, set[str]] = {}
    present: dict[str, int] = {}
    for record in records:
        for key, value in record.items():
            kind = "null" if value is None else names.get(type(value), "string")
            properties.setdefault(key, set()).add(kind)
            if value is not None:
                present[key] = present.get(key, 0) + 1
    schema_properties = {}
    for key, kinds in properties.items():
        if kinds >= {"integer", "number"}:
            kinds = kinds - {"integer"}
        ordered = sorted(kinds)
        schema_properties[key] = {"type": ordered[0] if len(ordered) == 1 else ordered}
    return {"$schema": "https://json-schema.org/draft/2020-12/schema", "type": "array",
            "items": {"type": "object", "properties": schema_properties,
                      "required": [key for key in properties if records and present.get(key) == len(records)]}}


def validate_data(source: Path, schema: Path | dict[str, Any], sheet: str | None = None,
                  limit: int = 200) -> list[dict[str, Any]]:
    """Check a data file against a JSON Schema. Returns problems; an empty list means valid.

    A schema of ``type: object`` is applied to every record; any other schema
    is applied to the whole list of records.
    """
    from jsonschema import Draft202012Validator
    from jsonschema.exceptions import SchemaError
    if isinstance(schema, Path):
        try:
            schema = json.loads(_read_text(schema)) if schema.suffix.lower() == ".json" \
                else __import__("yaml").safe_load(_read_text(schema))
        except Exception as exc:
            raise ValueError(f"The schema could not be read: {exc}")
    if not isinstance(schema, dict):
        raise ValueError("A schema must be a JSON object")
    try:
        Draft202012Validator.check_schema(schema)
    except SchemaError as exc:
        raise ValueError(f"The schema is not valid JSON Schema: {exc.message}")
    records = json.loads(json.dumps(load_records(source, sheet), default=str))
    validator = Draft202012Validator(schema)
    problems = []
    per_record = schema.get("type") == "object"
    for index, target in enumerate(records if per_record else [records]):
        for error in validator.iter_errors(target):
            path = list(error.absolute_path)
            record = index if per_record else (path[0] if path and isinstance(path[0], int) else None)
            field = ".".join(str(part) for part in (path if per_record else path[1:]))
            problems.append({"record": None if record is None else record + 1, "field": field, "message": error.message})
            if len(problems) >= limit:
                return problems
    return problems


# -- workbooks -----------------------------------------------------------

def export_sheets(source: Path, output_dir: Path, fmt: str = "CSV", overwrite: bool = False,
                  cancel_check: Callable[[], bool] | None = None) -> list[Path]:
    """Write every worksheet of a workbook to its own file. Empty sheets are skipped."""
    fmt = fmt.upper()
    if fmt not in DATA_WRITE_FORMATS or fmt == "XLSX":
        raise ValueError("Export sheets as CSV, TSV, JSON, JSONL, XML, or YAML")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Multi-sheet export reads .xlsx and .xlsm workbooks. "
                         "Convert other spreadsheets to XLSX first.")
    if not source.is_file():
        raise FileNotFoundError("The selected workbook no longer exists")
    from openpyxl import load_workbook
    book = load_workbook(source, read_only=True, data_only=True)
    outputs = []
    try:
        for sheet in book.worksheets:
            check_cancel(cancel_check)
            records = _sheet_records(sheet)
            if not records:
                continue
            safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", sheet.title).strip(" .") or "sheet"
            with StagedOutput(output_dir / f"{source.stem}-{safe}{DATA_WRITE_FORMATS[fmt]}", overwrite) as stage:
                save_records(records, stage.path, fmt)
            outputs.append(stage.output)
    finally:
        book.close()
    if not outputs:
        raise ValueError("This workbook has no sheets with data")
    return outputs
