"""Exercise real bundled engines without starting a GUI or downloading anything."""
from pathlib import Path
import json
import tempfile
import traceback


def run(destination: Path) -> int:
    checks = []
    with tempfile.TemporaryDirectory(prefix="shadow-studio-smoke-") as td:
        root = Path(td)
        def check(name, action):
            try:
                action()
                checks.append({"name": name, "passed": True})
            except Exception as exc:
                checks.append({"name": name, "passed": False, "error": str(exc), "traceback": traceback.format_exc()})

        def office():
            from docx import Document
            from pptx import Presentation
            from openpyxl import Workbook
            from doc_converter import convert_document
            from pypdf import PdfReader
            doc = Document()
            doc.add_paragraph("Agency document smoke test")
            doc.save(root / "proposal.docx")
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Agency pitch"
            deck.save(root / "pitch.pptx")
            book = Workbook()
            book.active.append(["Budget", 500])
            book.save(root / "budget.xlsx")
            book.close()
            for source in (root / "proposal.docx", root / "pitch.pptx", root / "budget.xlsx"):
                result = convert_document(source, root / "converted", "PDF")
                if result.status != "Completed":
                    raise RuntimeError(result.error)
                if len(PdfReader(result.output_path).pages) < 1:
                    raise RuntimeError("Office export has no pages")

        def pdf():
            from pdf_tools import render_pdf_page, process_pdf
            from pypdf import PdfWriter
            source = root / "preview.pdf"
            with PdfWriter() as writer:
                writer.add_blank_page(width=300, height=400)
                writer.write(source)
            image, count = render_pdf_page(source)
            image.close()
            if count != 1:
                raise RuntimeError("Wrong PDF page count")
            process_pdf([source], root / "pages.zip", "images")

        def speech():
            from speech_engine import synthesize_speech, transcribe_media
            audio = synthesize_speech("Hello. This is the agency studio. We create marketing campaigns.", root / "voice.wav")
            result = transcribe_media(audio, root, "SRT")
            if result.status != "Completed":
                raise RuntimeError(result.error)
            if "agency" not in result.output_path.read_text(encoding="utf-8").lower():
                raise RuntimeError("Speech round trip did not recover the expected text")

        def audio():
            import wave
            from audio_tools import process_audio
            source = root / "silent.wav"
            with wave.open(str(source), "wb") as wav:
                wav.setnchannels(1)
                wav.setsampwidth(2)
                wav.setframerate(16000)
                wav.writeframes(b"\x00\x00" * 32000)
            process_audio(source, root / "delivery.mp3", "MP3", denoise=True)

        check("Office DOCX/PPTX/XLSX to PDF", office)
        check("PDF rendering and page image export", pdf)
        check("English speech synthesis and transcription", speech)
        check("Audio normalization, denoising, and MP3 export", audio)
        def expanded_documents():
            from studio_actions import run_action
            from reportlab.pdfgen import canvas
            source = root / "confidential.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.drawString(50, 700, "Secret client Acme")
            pdf.save()
            protected = run_action("pdf-protect", [source], root, options={"open_password": "smoke-secret"}).outputs[0]
            run_action("pdf-unlock", [protected], root, options={"password": "smoke-secret"})
            run_action("pdf-redact", [source], root, options={"terms": ["Acme"]})
            run_action("pdf-repair", [source], root)
            identity = run_action("pdf-identity", [], root, options={"name": "Smoke verification", "passphrase": "smoke-pass"}).outputs[0]
            signed = run_action("pdf-sign", [source], root, options={"identity": str(identity), "passphrase": "smoke-pass"}).outputs[0]
            if not run_action("pdf-signatures", [signed], root).ok:
                raise RuntimeError("Signed bytes did not verify")

        def data_delivery():
            from studio_actions import run_action
            source = root / "leads.csv"
            source.write_text("Name,Budget\nAcme,100\nAcme,100\n", encoding="utf-8")
            cleaned = run_action("data-clean", [source], root, fmt="XLSX", drop_duplicates=True).outputs[0]
            run_action("data-export-sheets", [cleaned], root)
            schema = run_action("data-schema", [source], root).outputs[0]
            if not run_action("data-validate", [source], root, options={"schema": str(schema)}).ok:
                raise RuntimeError("Inferred schema did not validate")
            package = run_action("delivery-create", [cleaned], root, fmt="7Z", options={"password": "smoke-package"}).outputs[0]
            if not run_action("delivery-verify", [package], root, options={"password": "smoke-package"}).ok:
                raise RuntimeError("Delivery checksums did not verify")

        def ocr():
            from PIL import Image, ImageDraw, ImageFont
            from studio_runtime import resource_root
            import ocr_engine
            fonts = list((resource_root() / "assets" / "fonts").glob("*.ttf"))
            image = Image.new("RGB", (1200, 350), "white")
            font = ImageFont.truetype(str(fonts[0]), 48) if fonts else ImageFont.load_default(size=48)
            ImageDraw.Draw(image).text((80, 80), "Agency invoice 2026", font=font, fill="black")
            source = root / "scan.png"
            image.save(source)
            text = ocr_engine.ocr_document(source, root, "TXT").read_text(encoding="utf-8")
            if "invoice" not in text.lower():
                raise RuntimeError("OCR did not recover the invoice text")
            ocr_engine.ocr_document(source, root, "PDF")
            # Each staged script must at least initialize its recognition session.
            for language in ocr_engine.available_ocr_languages():
                ocr_engine._engine(language)
                ocr_engine._engines.clear()

        def speakers():
            from studio_actions import run_action
            from studio_runtime import model_directory
            result = run_action("speech-batch", [root / "voice.wav"], root,
                                options={"format": "JSON", "model_path": str(model_directory() / "whisper-base"),
                                         "speakers": True, "speaker_count": 1})
            if not result.ok:
                raise RuntimeError(str(result.details))
            run_action("audio-isolate", [root / "voice.wav"], root)

        def translation():
            from translation_engine import installed_pairs, translate_texts
            pairs = installed_pairs()
            if not pairs:
                raise RuntimeError("No offline translation pairs installed")
            for source, target in pairs:
                sample = "Hello world." if source == "en" else "Hola mundo." if source == "es" else "Bonjour le monde." if source == "fr" else "Hello world."
                result = translate_texts([sample], source, target)
                if not result or not result[0].strip():
                    raise RuntimeError(f"Translation {source}-{target} returned no text")

        def print_delivery():
            from studio_actions import run_action
            eps = root / "logo.eps"
            eps.write_text("%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 100 100\n1 0 0 setrgbcolor\n0 0 100 100 rectfill\nshowpage\n")
            pdf = run_action("print-import", [eps], root).outputs[0]
            run_action("print-fonts", [pdf], root)
            run_action("print-preflight", [pdf], root)
            run_action("print-cmyk", [pdf], root)
            from PIL import Image
            source = root / "print-image.png"
            Image.new("RGB", (100, 100), "red").save(source)
            run_action("print-cmyk", [source], root)

        def projects():
            from agency_projects import run_project
            from archive_tools import verify_package
            from PIL import Image
            source = root / "social.png"
            Image.new("RGB", (1000, 800), "blue").save(source)
            result = run_project({"client": "Smoke", "project": "Launch", "preset": "Social square"}, [source], root / "delivery")
            if not result.ok or not verify_package(result.outputs[-1])["ok"]:
                raise RuntimeError("Project delivery failed")
            with Image.open(result.outputs[0]) as image:
                if image.size != (1080, 1080):
                    raise RuntimeError("Social square preset did not produce 1080x1080")

        check("Advanced PDF passwords, redaction, repair, and signatures", expanded_documents)
        check("Structured data, schemas, sheets, encrypted archives, and delivery integrity", data_delivery)
        check("OCR text, searchable PDF, and staged OCR script sessions", ocr)
        check("Multilingual transcription, speaker labels, and speech isolation", speakers)
        check("Installed offline translation pairs", translation)
        check("EPS import, print preflight, and ICC CMYK delivery", print_delivery)
        check("Project presets, naming, and delivery package", projects)
    destination.parent.mkdir(parents=True, exist_ok=True)
    passed = all(check["passed"] for check in checks)
    destination.write_text(json.dumps({"passed": passed, "checks": checks}, indent=2), encoding="utf-8")
    return 0 if passed else 1
